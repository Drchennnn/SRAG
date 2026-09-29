"""PPTX 解析和预处理包的小型集成测试。"""

import base64
import json
from pathlib import Path

import pytest

from app.rag.slide_chunker import SlideChunker
from app.rag.slide_preprocessor import SlideDeckPreprocessor
from app.schemas.document import SlideElementKind

pytest.importorskip("pptx")
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB"
    "9Y9Z4QAAAABJRU5ErkJggg=="
)


def _make_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Linear Regression"
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(4), Inches(1))
    text_box.text = "Minimize mean squared error."
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "feature"
    table.cell(0, 1).text = "weight"
    table.cell(1, 0).text = "x1"
    table.cell(1, 1).text = "0.5"
    image_path = path.with_suffix(".png")
    image_path.write_bytes(_PNG)
    slide.shapes.add_picture(str(image_path), Inches(7), Inches(5), Inches(0.5))
    presentation.save(path)


def test_pptx_preprocessor_exports_portable_deck_and_chunks(tmp_path: Path) -> None:
    pptx_path = tmp_path / "lecture.pptx"
    _make_pptx(pptx_path)
    preprocessor = SlideDeckPreprocessor(
        output_root=tmp_path / "output",
        chunker=SlideChunker(max_tokens=200, overlap_tokens=20),
        render_pptx_slides=False,
    )

    result = preprocessor.preprocess(pptx_path)

    assert len(result.deck.slides) == 1
    kinds = {element.kind for element in result.deck.slides[0].elements}
    assert {SlideElementKind.TITLE, SlideElementKind.TEXT, SlideElementKind.TABLE} <= kinds
    assert result.deck.slides[0].assets
    assert result.chunking.chunks[0].slide_no == 1
    assert result.deck_path.exists()
    assert result.chunks_path.exists()
    payload = json.loads(result.deck_path.read_text(encoding="utf-8"))
    assert payload["slides"][0]["assets"][0]["path"].startswith("assets/")
